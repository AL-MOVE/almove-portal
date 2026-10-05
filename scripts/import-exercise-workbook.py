#!/usr/bin/env python3
"""Convert the AL MOVE exercise workbook into the static application catalog."""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl


SHEET_NAME = "Exercicios"
REQUIRED_COLUMNS = (
    "Exercise_ID",
    "Exercicio",
    "Nome_Alternativo",
    "Exercicio_Base",
    "Variacao",
    "Padrao_Movimento",
    "Grupo_Muscular_Primario",
    "Musculos_Secundarios",
    "Cadeia_Cinetica",
    "Tipo_Contracao",
    "Unilateral_Bilateral",
    "Equipamento",
    "Nivel",
    "Plano_Movimento",
    "Categoria_Treino",
    "URL_Imagem",
    "URL_Video",
    "Notas_Tecnicas",
    "Contraindicacoes",
    "Ativo",
)


def text(value: object) -> str:
    return str(value or "").strip()


def normalized_key(value: object) -> str:
    ascii_value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", ascii_value.lower()).strip()


def string_list(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[,;]", text(value)) if part.strip()]


def is_active(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return text(value).lower() not in {"", "0", "false", "nao", "não", "inativo"}


def load_catalog(source: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Folha obrigatória em falta: {SHEET_NAME}")

    rows = workbook[SHEET_NAME].iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing_columns:
        raise ValueError("Colunas obrigatórias em falta: " + ", ".join(missing_columns))

    catalog: list[dict[str, object]] = []
    seen_ids: set[str] = set()
    seen_names: set[str] = set()
    for row_number, values in enumerate(rows, start=2):
        source_row = dict(zip(headers, values))
        if not any(value not in (None, "") for value in values):
            continue

        exercise_id = text(source_row["Exercise_ID"])
        name = text(source_row["Exercicio"])
        if not exercise_id or not name:
            raise ValueError(f"Linha {row_number}: ID e nome são obrigatórios")

        id_key = normalized_key(exercise_id)
        name_key = normalized_key(name)
        if id_key in seen_ids:
            raise ValueError(f"Linha {row_number}: ID duplicado: {exercise_id}")
        if name_key in seen_names:
            raise ValueError(f"Linha {row_number}: nome duplicado: {name}")
        seen_ids.add(id_key)
        seen_names.add(name_key)

        primary_group = text(source_row["Grupo_Muscular_Primario"])
        catalog.append(
            {
                "id": exercise_id,
                "nome": name,
                "nomeAlternativo": text(source_row["Nome_Alternativo"]),
                "exercicioBase": text(source_row["Exercicio_Base"]),
                "variacao": text(source_row["Variacao"]),
                "padraoMovimento": text(source_row["Padrao_Movimento"]),
                "grupoMuscular": primary_group,
                "musculosPrincipais": [primary_group] if primary_group else [],
                "musculosSecundarios": string_list(source_row["Musculos_Secundarios"]),
                "cadeiaCinetica": text(source_row["Cadeia_Cinetica"]),
                "tipoContracao": text(source_row["Tipo_Contracao"]),
                "lateralidade": text(source_row["Unilateral_Bilateral"]),
                "equipamento": text(source_row["Equipamento"]),
                "nivel": text(source_row["Nivel"]),
                "planoMovimento": text(source_row["Plano_Movimento"]),
                "categoriaTreino": text(source_row["Categoria_Treino"]),
                "urlImagem": text(source_row["URL_Imagem"]),
                "urlVideo": text(source_row["URL_Video"]),
                "instrucoes": text(source_row["Notas_Tecnicas"]),
                "contraindicacoes": text(source_row["Contraindicacoes"]),
                "ativo": is_active(source_row["Ativo"]),
                "origem": "catalogo-al-move-v3",
            }
        )

    return catalog


def write_module(catalog: list[dict[str, object]], destination: Path) -> None:
    payload = json.dumps(catalog, ensure_ascii=False, indent=2)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        "// Gerado por scripts/import-exercise-workbook.py. Não editar manualmente.\n"
        f"export const ALMOVE_EXERCISE_CATALOG = Object.freeze({payload});\n",
        encoding="utf-8",
        newline="\n",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    catalog = load_catalog(args.source)
    if len(catalog) < 100:
        raise ValueError(f"Catálogo inesperadamente pequeno: {len(catalog)} exercícios")
    write_module(catalog, args.destination)
    print(f"{len(catalog)} exercícios exportados para {args.destination}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"Erro: {error}", file=sys.stderr)
        raise
